import base64
import os
import struct
import sys
import unittest
import zlib
from io import BytesIO

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import docx_writer  # noqa: E402


def _png(w=4, h=2):
    def chunk(t, d):
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * w for _ in range(h))
    return base64.b64encode(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")).decode()


class TestDocx(unittest.TestCase):
    def test_se_abre_con_python_docx(self):
        try:
            import docx
        except ImportError:
            self.skipTest("python-docx no instalado")
        data = docx_writer.build([
            {"t": "h", "n": 1, "text": "Memoria de cálculo — ñandú"}, {"t": "p", "text": "Hola & <mundo>", "b": True, "color": "#1a7f4b"},
            {"t": "ul", "items": ["uno", "dos"]},
            {"t": "table", "w": [3, 1, 1], "rows": [[{"text": "A", "b": True, "bg": "#eef2f7"}, "B", {"text": "C", "al": "right"}],
                                                     [{"text": "grupo", "span": 3, "bg": "#e8f0fb"}], ["x", "y", "z"]]},
            {"t": "imgs", "items": [{"png": _png()}, {"png": _png(6, 3)}], "cm": 16}, {"t": "pb"}, {"t": "h", "n": 2, "text": "Segunda"}])
        d = docx.Document(BytesIO(data))
        self.assertEqual(d.paragraphs[0].text, "Memoria de cálculo — ñandú")
        self.assertEqual(d.paragraphs[0].style.name, "Heading 1")
        self.assertEqual(len(d.tables), 1)
        self.assertEqual(d.tables[0].rows[2].cells[1].text, "y")
        self.assertEqual(len(d.inline_shapes), 2)
        self.assertIn("Hola & <mundo>", [p.text for p in d.paragraphs])

    def test_png_invalido(self):
        with self.assertRaises(ValueError):
            docx_writer.build([{"t": "imgs", "items": [{"png": base64.b64encode(b"no es png").decode()}]}])


if __name__ == "__main__":
    unittest.main()
