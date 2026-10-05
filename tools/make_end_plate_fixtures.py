#!/usr/bin/env python3
"""Genera tests/fixture_end_plate.json: variantes de la hoja de Excel recalculadas con LibreOffice.

Cada variante cambia entradas de DISEÑO, se recalcula con `soffice --headless` y se guardan los valores
resultantes (CALCULO filas 111-138 y ratios/demandas/capacidades de la tabla de verificaciones).
Uso:  python3 tools/make_end_plate_fixtures.py "END_PLATE_AISC_DG4 (2).xlsx"
"""
import copy
import json
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SPEC = json.load(open(os.path.join(ROOT, "end_plate", "inputs_spec.json"), encoding="utf-8"))
ROW = {c["name"]: c["row"] for s in SPEC["secciones"] for c in s["campos"]}
COL = {c["name"]: ("E" if c["name"].startswith("DIS_E") else "C") for s in SPEC["secciones"] for c in s["campos"]}

BASE_COMBOS = [("1.2D+1.6L", 9.5, 6), ("1.2D+1.0E+L", 14, 7), ("0.9D+1.0E", 11.5, 4.5), ("1.2D+1.0W+L", 8, 5)]
VARIANTES = {
    "base_4E": {},
    "4ES": {"cfg": "4ES"},
    "8ES": {"cfg": "8ES", "pl_pb": 90, "pl_tp": 28},
    "4E_sismico": {"modo": "Sísmico (AISC 358)"},
    "4ES_sismico": {"cfg": "4ES", "modo": "Sísmico (AISC 358)", "st_t": 12},
    "8ES_sismico": {"cfg": "8ES", "modo": "Sísmico (AISC 358)", "pl_pb": 90, "pl_tp": 32, "sis_L": 9},
    "sin_rigidizadores_dos_lados": {"cp_usar": "No", "dos_lados": "Sí", "Mu_op": 6, "Pu_col": 80, "co_dtop": 300},
    "armado_4ES": {"cfg": "4ES", "vg_perfil": "ARMADO (flejes soldados)", "DIS_E15": 550, "DIS_E16": 220,
                   "DIS_E17": 8, "DIS_E18": 14, "co_perfil": "ARMADO (flejes soldados)", "DIS_E25": 450,
                   "DIS_E26": 300, "DIS_E27": 9, "DIS_E28": 18, "pl_bp": 230, "pl_g": 130, "pn_diam": "1\"",
                   "pn_grado": "A490-N (roscas incl.)", "pl_tp": 30},
}


def main(xlsx):
    out = {}
    tmp = tempfile.mkdtemp()
    for nombre, cambios in VARIANTES.items():
        wb = openpyxl.load_workbook(xlsx)
        ws = wb["DISEÑO"]
        for k, v in cambios.items():
            ws[f"{COL[k]}{ROW[k]}"] = v
        for i, (n, m, v) in enumerate(BASE_COMBOS):
            ws[f"G{23+i}"], ws[f"H{23+i}"], ws[f"I{23+i}"] = n, m, v
        src = os.path.join(tmp, f"{nombre}.xlsx")
        wb.save(src)
        outdir = os.path.join(tmp, "out")
        os.makedirs(outdir, exist_ok=True)
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", outdir, src],
                       check=True, capture_output=True, timeout=180)
        wv = openpyxl.load_workbook(os.path.join(outdir, f"{nombre}.xlsx"), data_only=True)
        cal, dis = wv["CALCULO"], wv["DISEÑO"]
        rows = {}
        for r in range(111, 139):
            vals = [cal.cell(row=r, column=4 + j).value for j in range(4)]
            if all(isinstance(x, (int, float)) for x in vals):
                rows[str(r)] = vals
        checks = {}
        for r in range(12, 45):
            p, q, s = (dis.cell(row=r, column=c).value for c in (16, 17, 19))
            if s is not None:
                checks[str(r)] = [p, q, s]
        g = {str(r): cal.cell(row=r, column=4).value for r in range(5, 107)
             if isinstance(cal.cell(row=r, column=4).value, (int, float))}
        out[nombre] = {"cambios": cambios, "combos": [list(c) for c in BASE_COMBOS], "rows": rows,
                       "checks": checks, "global": g}
        print(nombre, "ok", len(rows), len(checks))
    json.dump(out, open(os.path.join(ROOT, "tests", "fixture_end_plate.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=0)
    shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main(sys.argv[1])
