#!/usr/bin/env python3
"""Genera tests/fixtures/<módulo>.json: variantes de la hoja de Excel recalculadas con LibreOffice Calc.

Cada variante cambia entradas de DISEÑO, se recalcula con `soffice --headless` y se guardan los valores
resultantes (CALCULO y tabla de verificaciones). Uso:
    python3 tools/make_fixtures.py end_plate "END_PLATE_AISC_DG4 (2).xlsx"
    python3 tools/make_fixtures.py bfp "BFP_AISC358 (2).xlsx"
    python3 tools/make_fixtures.py rodilla "RODILLA_CUMBRERA_GALPON (2).xlsx"
"""
import importlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

VARIANTES = {
    "end_plate": {
        "base_4E": {},
        "4ES": {"cfg": "4ES"},
        "8ES": {"cfg": "8ES", "pl_pb": 90, "pl_tp": 28},
        "4E_sismico": {"modo": "Sísmico (AISC 358)"},
        "4ES_sismico": {"cfg": "4ES", "modo": "Sísmico (AISC 358)", "st_t": 12},
        "8ES_sismico": {"cfg": "8ES", "modo": "Sísmico (AISC 358)", "pl_pb": 90, "pl_tp": 32, "sis_L": 9},
        "sin_rigidizadores_dos_lados": {"cp_usar": "No", "dos_lados": "Sí", "Mu_op": 6, "Pu_col": 80, "co_dtop": 300},
        "armado_4ES": {"cfg": "4ES", "vg_perfil": "ARMADO (flejes soldados)", "DIS_E15": 550, "DIS_E16": 220,
                       "DIS_E17": 8, "DIS_E18": 14, "co_perfil": "ARMADO (flejes soldados)", "DIS_E25": 450,
                       "DIS_E26": 300, "DIS_E27": 9, "DIS_E28": 18, "pl_bp": 230, "pl_g": 130, "pn_diam": "1\"",
                       "pn_grado": "A490-N (roscas incl.)", "pl_tp": 30},
    },
    "bfp": {
        "base_IMF": {},
        "SMF": {"sistema": "SMF"},
        "armado_10_pernos": {"vg_perfil": "ARMADO (flejes soldados)", "DIS_E13": 500, "DIS_E14": 200, "DIS_E15": 8,
                             "DIS_E16": 14, "n_b": 10, "pn_diam": "1\"", "pl_tp": 28, "pl_bfp": 220},
        "perfil_menor_A325": {"vg_perfil": "VK300X150X4X10", "pn_grado": "A325-N (roscas incl.)", "n_b": 6,
                              "pl_tp": 20, "pl_bfp": 180, "L_m": 6},
        "sin_rigidizadores_dos_lados": {"cp_usar": "No", "dos_lados": "Sí", "Pu_col": 60, "co_dtop": 400},
        "alma_distinta": {"w_n": 5, "w_diam": "7/8\"", "w_t": 12, "w_w": 10, "setback": 20},
    },
    "cortante_vv": {
        "base_placa_convencional": {},
        "placa_extendida": {"tipo": "Placa simple extendida", "pl_tp": 10, "pl_a": 120, "pl_Leh": 50, "Ru": 8},
        "doble_angulo": {"tipo": "Doble ángulo apernado", "n_b": 4, "Ru": 12, "Ru_op": 6, "pn_diam": "7/8\"", "an_t": 10},
        "doble_angulo_sin_destaje": {"tipo": "Doble ángulo apernado", "destaje": "Sin destaje", "Ru": 9, "an_lb": 89, "an_ls": 89},
        "destaje_doble": {"destaje": "Destaje doble (sup. = inf.)", "cope_dc": 25, "cope_c": 100, "Ru": 5},
        "armado_A490_ranura": {"vg_perfil": "ARMADO (flejes soldados)", "DIS_E16": 360, "DIS_E17": 170, "DIS_E18": 7,
                               "DIS_E19": 12, "pn_grado": "A490-N (roscas incl.)", "agujero": "SSLT (ranura corta horiz.)",
                               "n_b": 5, "pn_s": 80, "Ru": 14, "pl_tp": 10, "pl_acero": "A572 Gr50", "destaje": "Sin destaje"},
    },
    "rodilla": {
        "base_rodilla": {},
        "cumbrera": {"tipo": "Cumbrera (placa a placa)"},
        "ras_2filas_sup": {"cfg_u": "Al ras — 2 filas (4 pernos)", "cfg_d": "Extendida — 1+1 (4E)"},
        "cjp_inclinada": {"sd_tipo": "CJP", "theta": 15, "vg_dh": 0},
        "sin_rigid_doubler": {"cp_usar": "No", "dp_t": 10, "Pu_col": 25},
        "armado_ras": {"cfg_u": "Al ras — 1 fila (2 pernos)", "vg_perfil": "ARMADO (flejes soldados)", "DIS_E18": 450,
                       "DIS_E19": 200, "DIS_E20": 8, "DIS_E21": 14, "pn_diam": "1\"", "pl_tp": 28},
    },
}


def main(modulo, xlsx):
    model = importlib.import_module(f"{modulo}.xl_model")
    spec = json.load(open(os.path.join(ROOT, modulo, "inputs_spec.json"), encoding="utf-8"))
    cell = {c["name"]: ("E" if c["name"].startswith("DIS_E") else "C") + str(c["row"])
            for s in spec["secciones"] for c in s["campos"]}
    wb0 = openpyxl.load_workbook(xlsx)
    cbnames = {}
    for n, d in wb0.defined_names.items():
        if n.startswith("cb_"):
            m = re.match(r"'?DISEÑO'?!\$([A-Z]+)\$(\d+):\$[A-Z]+\$(\d+)", d.attr_text)
            cbnames[n] = (m.group(1), int(m.group(2)))
    combos = spec["combos"]
    ncomb = 4 if combos else 0
    combo_rows = [r for r in range(1, wb0["CALCULO"].max_row + 1)
                  if isinstance(wb0["CALCULO"].cell(row=r, column=5).value, str)
                  and wb0["CALCULO"].cell(row=r, column=5).value.startswith("=")]
    out = {}
    tmp = tempfile.mkdtemp()
    for nombre, cambios in VARIANTES[modulo].items():
        wb = openpyxl.load_workbook(xlsx)
        ws = wb["DISEÑO"]
        for k, v in cambios.items():
            ws[cell[k]] = v
        cargas = []
        if combos:
            cargas = combos["filas"][:ncomb]
            for i, fila in enumerate(cargas):
                for col in combos["cols"]:
                    L, r0 = cbnames[col["name"]]
                    ws[f"{L}{r0 + i}"] = fila[col["key"]]
        src = os.path.join(tmp, f"{nombre}.xlsx")
        wb.save(src)
        outdir = os.path.join(tmp, "out")
        os.makedirs(outdir, exist_ok=True)
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", outdir, src],
                       check=True, capture_output=True, timeout=300)
        wv = openpyxl.load_workbook(os.path.join(outdir, f"{nombre}.xlsx"), data_only=True)
        cal, dis = wv["CALCULO"], wv["DISEÑO"]

        def val(x):
            return "ERR" if isinstance(x, str) and (x.startswith("#") or x.startswith("Err:")) else x

        glob = {str(r): val(cal.cell(row=r, column=4).value) for r in range(1, cal.max_row + 1)
                if r not in combo_rows and cal.cell(row=r, column=4).value is not None
                and not isinstance(cal.cell(row=r, column=4).value, str) or
                (r not in combo_rows and isinstance(cal.cell(row=r, column=4).value, str)
                 and cal.cell(row=r, column=4).value.startswith(("#", "Err:")))}
        rows = {str(r): [val(cal.cell(row=r, column=4 + j).value) for j in range(ncomb)] for r in combo_rows}
        checks = {}
        for (row, tipo, *_rest) in model.CHECKS:
            if tipo != "CHEQUEO":
                continue
            checks[str(row)] = [val(dis[f"{c}{row}"].value) for c in (model.DEM_COL, model.CAP_COL, model.RATIO_COL)]
        out[nombre] = {"cambios": cambios, "combos": cargas, "global": glob, "rows": rows, "checks": checks}
        print(modulo, nombre, "ok", len(glob), len(rows), len(checks))
    os.makedirs(os.path.join(ROOT, "tests", "fixtures"), exist_ok=True)
    json.dump(out, open(os.path.join(ROOT, "tests", "fixtures", f"{modulo}.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=0)
    shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
