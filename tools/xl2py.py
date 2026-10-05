#!/usr/bin/env python3
"""Traduce las fórmulas de una hoja de cálculo de conexión (DISEÑO / CALCULO / CATALOGOS) a un módulo Python.

Uso:  python3 tools/xl2py.py "END_PLATE_AISC_DG4 (2).xlsx" end_plate

Genera en el paquete de destino:
    xl_model.py        filas de fórmulas (expresiones Python con comentarios) — NO editar a mano
    catalogos.json     columnas de la hoja CATALOGOS
    inputs_spec.json   campos de entrada (etiqueta, unidad, valor por defecto, opciones) para la interfaz

Cada fila de la hoja CALCULO se convierte en una expresión que lleva su símbolo, descripción y unidad,
de modo que la memoria de cálculo del programa es la misma de la hoja de Excel.
Requiere openpyxl (sólo para regenerar; el programa en uso no lo necesita).
"""
import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

TOKEN = re.compile(r'''\s*(?:
 (?P<str>"(?:[^"]|"")*")|
 (?P<num>\d+\.?\d*(?:[eE][+-]?\d+)?|\.\d+)|
 (?P<ref>(?:(?:'[^']+'|[A-Za-z_0-9ÁÉÍÓÚÑáéíóúñ]+)!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)(?![A-Za-z0-9_(])|
 (?P<name>[A-Za-z_][A-Za-z0-9_.]*)|
 (?P<op><=|>=|<>|[-+*/^&=<>(),%])
)''', re.X)


def tokenize(s):
    pos, out = 0, []
    while pos < len(s):
        if s[pos:].strip() == "":
            break
        m = TOKEN.match(s, pos)
        if not m:
            raise ValueError(f"token no reconocido en {s[pos:pos+30]!r} de {s!r}")
        pos = m.end()
        kind = m.lastgroup
        out.append((kind, m.group(kind)))
    return out


class Translator:
    """Fórmula de Excel → expresión Python. `ctx` decide cómo se llaman las referencias a celdas."""

    def __init__(self, names, sheet, col_letter=None, cell_name=None):
        self.names = names              # nombre definido → (hoja, celda)
        self.sheet = sheet              # hoja de la fórmula
        self.col = col_letter           # columna (para filas por combinación)
        self.cell_name = cell_name      # (hoja, celda) → identificador Python

    # -- análisis sintáctico ---------------------------------------------------------------
    def tr(self, formula):
        self.toks = tokenize(formula.lstrip("="))
        self.i = 0
        code = self.cmp()
        if self.i != len(self.toks):
            raise ValueError(f"sobran tokens en {formula!r}")
        return code

    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else (None, None)

    def eat(self, val=None):
        t = self.toks[self.i]
        self.i += 1
        if val is not None and t[1] != val:
            raise ValueError(f"se esperaba {val!r}, llegó {t!r}")
        return t

    def cmp(self):
        a = self.cat()
        while self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.eat()[1]
            b = self.cat()
            a = {"=": f"_eq({a},{b})", "<>": f"(not _eq({a},{b}))"}.get(op, f"({a}{op}{b})")
        return a

    def cat(self):
        a = self.add()
        while self.peek()[1] == "&":
            self.eat()
            a = f"(_s({a})+_s({self.add()}))"
        return a

    def add(self):
        a = self.mul()
        while self.peek()[1] in ("+", "-"):
            op = self.eat()[1]
            a = f"({a}{op}{self.mul()})"
        return a

    def mul(self):
        a = self.pow()
        while self.peek()[1] in ("*", "/"):
            op = self.eat()[1]
            a = f"({a}{op}{self.pow()})"
        return a

    def pow(self):
        a = self.unary()
        while self.peek()[1] == "^":
            self.eat()
            a = f"({a}**{self.unary()})"
        return a

    def unary(self):
        if self.peek()[1] == "-":
            self.eat()
            return f"(-{self.unary()})"
        if self.peek()[1] == "+":
            self.eat()
            return self.unary()
        return self.primary()

    def primary(self):
        kind, val = self.eat()
        if kind == "num":
            return repr(float(val)) if ("." in val or "e" in val.lower()) else val
        if kind == "str":
            return repr(val[1:-1].replace('""', '"'))
        if kind == "op" and val == "(":
            e = self.cmp()
            self.eat(")")
            return f"({e})"
        if kind == "ref":
            return self.ref(val)
        if kind == "name":
            if self.peek()[1] == "(":
                self.eat("(")
                args = []
                if self.peek()[1] != ")":
                    args.append(self.cmp())
                    while self.peek()[1] == ",":
                        self.eat()
                        args.append(self.cmp())
                self.eat(")")
                return self.func(val.upper(), args)
            return self.name(val)
        raise ValueError(f"token inesperado {kind} {val}")

    # -- referencias -------------------------------------------------------------------------
    def name(self, n):
        if n == "IDX_":                                   # nº de la combinación en curso
            return n
        if n not in self.names:
            raise ValueError(f"nombre no definido: {n}")
        sh, cell = self.names[n]
        if ":" in cell:                                   # rango con nombre (cb_nom, cb_M, cb_V)
            return n
        return n

    def ref(self, text):
        sheet = self.sheet
        if "!" in text:
            sheet, text = text.split("!")
            sheet = sheet.strip("'")
        text = text.replace("$", "")
        if ":" in text:
            a, b = text.split(":")
            ca, ra = re.match(r"([A-Z]+)(\d+)", a).groups()
            cb, rb = re.match(r"([A-Z]+)(\d+)", b).groups()
            if sheet == "CATALOGOS":
                return f"_cat({ca!r},{ra},{cb!r},{rb})"
            if sheet == "CALCULO" and ra == rb:
                return f"_row({ra})"
            raise ValueError(f"rango no soportado {sheet}!{text}")
        col, row = re.match(r"([A-Z]+)(\d+)", text).groups()
        return self.cell_name(sheet, col, int(row), self.col)

    def func(self, f, a):
        if f == "IF":
            return f"(({a[1]}) if ({a[0]}) else ({a[2] if len(a) > 2 else 'False'}))"
        simple = {"MIN": "_min", "MAX": "_max", "ABS": "abs", "ISNUMBER": "_isnum", "N": "_n",
                  "MATCH": "_match", "INDEX": "_index", "COUNT": "_count", "AND": "_and", "OR": "_or",
                  "SQRT": "_sqrt", "TAN": "math.tan", "RADIANS": "math.radians", "TEXT": "_text"}
        if f in simple:
            return f"{simple[f]}({','.join(a)})"
        if f == "PI":
            return "math.pi"
        if f == "LEFT":
            return f"_s({a[0]})[:int({a[1]})]"
        if f == "IFERROR":
            return f"_iferror(lambda: {a[0]}, lambda: {a[1]})"
        raise ValueError(f"función no soportada: {f}")


def main(xlsx, outdir):
    wb = openpyxl.load_workbook(xlsx)
    wv = openpyxl.load_workbook(xlsx, data_only=True)
    dis, cal, cat = wb["DISEÑO"], wb["CALCULO"], wb["CATALOGOS"]
    names = {n: (d.attr_text.split("!")[0].strip("'"), d.attr_text.split("!")[1].replace("$", ""))
             for n, d in wb.defined_names.items()}
    cell2name = {}
    for n, (sh, cell) in names.items():
        if ":" not in cell:
            cell2name[(sh, cell)] = n

    def cell_name(sheet, col, row, combo_col):
        if sheet == "DISEÑO":
            return cell2name.get((sheet, f"{col}{row}"), f"DIS_{col}{row}")
        if sheet == "CALCULO":
            if row >= 109:
                return f"r{row}"
            return cell2name.get((sheet, f"{col}{row}"), f"g{row}")
        raise ValueError(sheet)

    os.makedirs(outdir, exist_ok=True)

    # ── catálogos ─────────────────────────────────────────────────────────────
    catd = {}
    for col in range(1, cat.max_column + 1):
        L = get_column_letter(col)
        vals = [cat.cell(row=r, column=col).value for r in range(1, cat.max_row + 1)]
        if any(v is not None for v in vals):
            catd[L] = vals
    json.dump(catd, open(os.path.join(outdir, "catalogos.json"), "w", encoding="utf-8"), ensure_ascii=False)

    # ── campos de entrada de DISEÑO ──────────────────────────────────────────────
    sections, cur = [], None
    validations = {}
    for dv in dis.data_validations.dataValidation:
        for rng in str(dv.sqref).split():
            validations[rng] = dv.formula1
    derived, inputs = [], []

    def opts(formula):
        m = re.match(r"CATALOGOS!\$([A-Z]+)\$(\d+):\$[A-Z]+\$(\d+)", formula or "")
        if not m:
            return None
        L, a, b = m.group(1), int(m.group(2)), int(m.group(3))
        return [v for v in catd[L][a - 1:b] if v is not None]

    for r in range(6, 67):
        b = dis.cell(row=r, column=2).value
        c = dis.cell(row=r, column=3)
        if b and c.value is None and r != 14 and not str(b).startswith("Celdas"):
            cur = {"titulo": str(b), "campos": []}
            sections.append(cur)
            continue
        if c.value is None or cur is None:
            continue
        nm = cell_name("DISEÑO", "C", r, None)
        unidad = dis.cell(row=r, column=4).value or ""
        if isinstance(c.value, str) and c.value.startswith("="):
            derived.append((r, nm, str(b), unidad, c.value))
            e = dis.cell(row=r, column=5)
            continue
        campo = {"name": nm, "label": str(b), "unit": unidad, "default": c.value, "row": r}
        o = opts(validations.get(f"C{r}"))
        if o:
            campo["options"] = o
        elif isinstance(c.value, str):
            campo["kind"] = "text"
        cur["campos"].append(campo)
        inputs.append(campo)
        # entradas "armado" en la columna E
        e = dis.cell(row=r, column=5)
        if e.value is not None and not (isinstance(e.value, str) and e.value.startswith("=")) \
                and dis.cell(row=r, column=5).value != "Armado (mm)" and isinstance(e.value, (int, float)):
            pass
    # columnas E (armado): filas 15-18 y 25-28
    for r in list(range(15, 19)) + list(range(25, 29)):
        e = dis.cell(row=r, column=5)
        nm = f"DIS_E{r}"
        campo = {"name": nm, "label": str(dis.cell(row=r, column=2).value) + " (armado)", "unit": "mm",
                 "default": e.value, "row": r, "armado_de": "vg_perfil" if r < 20 else "co_perfil"}
        for s in sections:
            if s["campos"] and s["campos"][0]["row"] < r <= (s["campos"][-1]["row"] + 1) or \
                    (s["campos"] and s["campos"][0]["row"] - 2 < r < s["campos"][0]["row"] + 12):
                pass
        inputs.append(campo)
    # asignar armados a su sección por rango de filas
    for campo in inputs:
        if campo["name"].startswith("DIS_E"):
            for s in sections:
                rows = [c["row"] for c in s["campos"] if not c["name"].startswith("DIS_E")]
                if rows and min(rows) <= campo["row"] <= max(rows):
                    s["campos"].append(campo)
                    break
    # tabla de combinaciones
    combos = []
    for r in range(23, 33):
        combos.append({"nombre": dis.cell(row=r, column=7).value, "M": dis.cell(row=r, column=8).value,
                       "V": dis.cell(row=r, column=9).value})
    spec = {"secciones": sections, "combos": combos}
    json.dump(spec, open(os.path.join(outdir, "inputs_spec.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)

    # ── traducción ───────────────────────────────────────────────────────────────
    def T(formula, sheet, col=None):
        return Translator(names, sheet, col, cell_name).tr(formula)

    out = ['"""Modelo generado por tools/xl2py.py — NO editar a mano (regenerar desde la hoja de Excel)."""', ""]
    out.append("DERIVED = [  # (fila, nombre, etiqueta, unidad, expresión)  — fórmulas de la hoja DISEÑO")
    for r, nm, lab, un, f in derived:
        out.append(f"    ({r}, {nm!r}, {lab!r}, {un!r}, {T(f, 'DISEÑO')!r}),")
    out.append("]\n")

    glob = []
    for r in range(5, 107):
        v = cal.cell(row=r, column=4).value
        a, b, c = (cal.cell(row=r, column=k).value for k in (1, 2, 3))
        if v is None:
            if a and not b:
                glob.append((r, "SECCION", str(a), "", None, None))
            continue
        nm = cell_name("CALCULO", "D", r, None)
        code = T(v, "CALCULO") if isinstance(v, str) and v.startswith("=") else repr(v)
        glob.append((r, str(a), str(b), str(c or ""), nm, code))
    out.append("GLOBAL = [  # (fila, símbolo, descripción, unidad, nombre, expresión)  — hoja CALCULO, sección 1-7")
    for g in glob:
        out.append(f"    {g!r},")
    out.append("]\n")

    # filas por combinación (rows 110-138): se traduce la columna D y se verifica contra E..M
    combo_rows = []
    for r in range(110, 139):
        v = cal.cell(row=r, column=4).value
        a, b, c = (cal.cell(row=r, column=k).value for k in (1, 2, 3))
        if not (isinstance(v, str) and v.startswith("=")):
            continue
        tpl = re.sub(r"INDEX\((cb_[A-Za-z]+),1\)", r"INDEX(\1,IDX)", v)
        tpl = re.sub(r"IF\(1=1,", "IF(IDX=1,", tpl)
        for j in range(2, 11):
            col = get_column_letter(3 + j)
            esperado = cal.cell(row=r, column=3 + j).value
            gen = tpl.replace("IDX", str(j))
            gen = re.sub(r"\bD(\d{3})\b", lambda m: f"{col}{m.group(1)}", gen)
            assert gen == esperado, (r, col, gen, esperado)
        code = T(tpl.replace("IDX", "IDX_"), "CALCULO", "D")
        combo_rows.append((r, str(a), str(b), str(c or ""), code))
    out.append("COMBO = [  # (fila, símbolo, descripción, unidad, expresión)  — hoja CALCULO, sección 8; IDX_ = nº de combinación")
    for g in combo_rows:
        out.append(f"    {g!r},")
    out.append("]\n")

    # ── tabla de verificaciones (DISEÑO N..S, filas 11-44) ────────────────────────────
    checks = []
    for r in range(11, 45):
        n = dis.cell(row=r, column=14).value
        if n is None:
            continue
        P = dis.cell(row=r, column=16).value
        if P is None:
            checks.append((r, "GRUPO", str(n), "", "", None, None, None, None))
            continue
        o = dis.cell(row=r, column=15).value
        un = dis.cell(row=r, column=18).value
        Q = dis.cell(row=r, column=17).value
        S = dis.cell(row=r, column=19).value
        U = dis.cell(row=r, column=21).value
        tr = lambda f: T(f, "DISEÑO") if isinstance(f, str) and f.startswith("=") else repr(f)
        checks.append((r, "CHEQUEO", str(n), str(o), str(un), tr(P), tr(Q), tr(S), tr(U) if U is not None else "'—'"))
    out.append("CHECKS = [  # (fila, tipo, nombre, referencia, unidad, demanda, capacidad, ratio, combo)")
    for g in checks:
        out.append(f"    {g!r},")
    out.append("]\n")
    with open(os.path.join(outdir, "xl_model.py"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(out))
    print(f"{len(derived)} derivadas · {len(glob)} globales · {len(combo_rows)} por combinación · {len(checks)} chequeos")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
